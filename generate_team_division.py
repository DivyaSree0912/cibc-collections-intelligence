from docx import Document
from docx.shared import Pt, RGBColor, Inches, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

# ── helpers ───────────────────────────────────────────────────────────────────

def shade_cell(cell, fill_hex):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), fill_hex)
    tcPr.append(shd)

def add_heading(doc, text, level=1, color='1E3A5F'):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = p.add_run(text)
    run.bold = True
    sizes = {1: 18, 2: 14, 3: 12, 4: 11}
    run.font.size = Pt(sizes.get(level, 11))
    r, g, b = tuple(int(color[i:i+2], 16) for i in (0, 2, 4))
    run.font.color.rgb = RGBColor(r, g, b)
    p.paragraph_format.space_before = Pt(14 if level == 1 else 10)
    p.paragraph_format.space_after = Pt(4)
    return p

def add_para(doc, text, bold=False, italic=False, size=10.5, color='222222',
             align=WD_ALIGN_PARAGRAPH.LEFT, space_before=2, space_after=5):
    p = doc.add_paragraph()
    p.alignment = align
    run = p.add_run(text)
    run.bold = bold
    run.italic = italic
    run.font.size = Pt(size)
    r, g, b = tuple(int(color[i:i+2], 16) for i in (0, 2, 4))
    run.font.color.rgb = RGBColor(r, g, b)
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.space_after = Pt(space_after)
    return p

def add_bullet(doc, text, color='374151', bold_prefix=None):
    p = doc.add_paragraph(style='List Bullet')
    if bold_prefix:
        rb = p.add_run(bold_prefix)
        rb.bold = True
        rb.font.size = Pt(10)
        rb.font.color.rgb = RGBColor(0x1E, 0x3A, 0x5F)
    run = p.add_run(text)
    run.font.size = Pt(10)
    r, g, b = tuple(int(color[i:i+2], 16) for i in (0, 2, 4))
    run.font.color.rgb = RGBColor(r, g, b)
    p.paragraph_format.left_indent = Inches(0.2)
    p.paragraph_format.space_after = Pt(3)
    return p

def add_divider(doc, color='2563EB'):
    p = doc.add_paragraph()
    pPr = p._p.get_or_add_pPr()
    pBdr = OxmlElement('w:pBdr')
    bottom = OxmlElement('w:bottom')
    bottom.set(qn('w:val'), 'single')
    bottom.set(qn('w:sz'), '6')
    bottom.set(qn('w:space'), '1')
    bottom.set(qn('w:color'), color)
    pBdr.append(bottom)
    pPr.append(pBdr)
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(8)

def add_callout(doc, label, text, bg='EFF6FF', bar='2563EB', label_color='1D4ED8'):
    table = doc.add_table(rows=1, cols=2)
    table.columns[0].width = Cm(0.45)
    table.columns[1].width = Cm(15.55)
    lc = table.cell(0, 0)
    rc = table.cell(0, 1)
    shade_cell(lc, bar)
    shade_cell(rc, bg)
    p = rc.paragraphs[0]
    run = p.add_run(label + '  ')
    run.bold = True
    run.font.size = Pt(9)
    r, g, b = tuple(int(label_color[i:i+2], 16) for i in (0, 2, 4))
    run.font.color.rgb = RGBColor(r, g, b)
    run2 = p.add_run(text)
    run2.font.size = Pt(9.5)
    run2.font.color.rgb = RGBColor(0x37, 0x41, 0x51)
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.left_indent = Inches(0.08)
    doc.add_paragraph().paragraph_format.space_after = Pt(5)

def add_code_block(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.3)
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(6)
    run = p.add_run(text)
    run.font.name = 'Courier New'
    run.font.size = Pt(8.5)
    run.font.color.rgb = RGBColor(0x1E, 0x3A, 0x5F)
    pPr = p._p.get_or_add_pPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), 'EEF2FF')
    pPr.append(shd)
    return p

def make_table(doc, headers, rows, header_bg='1E3A5F', row_bgs=None):
    tbl = doc.add_table(rows=len(rows)+1, cols=len(headers))
    tbl.style = 'Table Grid'
    for i, h in enumerate(headers):
        cell = tbl.cell(0, i)
        shade_cell(cell, header_bg)
        run = cell.paragraphs[0].add_run(h)
        run.bold = True
        run.font.size = Pt(9.5)
        run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
    for ri, row in enumerate(rows, start=1):
        for ci, val in enumerate(row):
            cell = tbl.cell(ri, ci)
            bg = 'F8FAFC' if ri % 2 == 0 else 'FFFFFF'
            if row_bgs and ci < len(row_bgs):
                bg = row_bgs[ci]
            shade_cell(cell, bg)
            p = cell.paragraphs[0]
            is_bold = ci == 0
            run = p.add_run(val)
            run.bold = is_bold
            run.font.size = Pt(9.5)
            run.font.color.rgb = RGBColor(0x1E, 0x3A, 0x5F) if is_bold else RGBColor(0x37, 0x41, 0x51)
    doc.add_paragraph().paragraph_format.space_after = Pt(4)
    return tbl

# ── document setup ─────────────────────────────────────────────────────────────

doc = Document()
section = doc.sections[0]
section.top_margin    = Cm(2.0)
section.bottom_margin = Cm(2.0)
section.left_margin   = Cm(2.5)
section.right_margin  = Cm(2.0)
style = doc.styles['Normal']
style.font.name = 'Calibri'
style.font.size = Pt(10.5)

# ═══════════════════════════════════════════════════════════════════════════════
# TITLE PAGE
# ═══════════════════════════════════════════════════════════════════════════════

tp = doc.add_paragraph()
tp.alignment = WD_ALIGN_PARAGRAPH.CENTER
tr = tp.add_run('Collections Intelligence Platform')
tr.bold = True
tr.font.size = Pt(26)
tr.font.color.rgb = RGBColor(0x1E, 0x3A, 0x5F)
tp.paragraph_format.space_before = Pt(8)
tp.paragraph_format.space_after = Pt(6)

sp = doc.add_paragraph()
sp.alignment = WD_ALIGN_PARAGRAPH.CENTER
sr = sp.add_run('Team Work Division: 2-Member Hackathon Execution Plan')
sr.italic = True
sr.font.size = Pt(14)
sr.font.color.rgb = RGBColor(0x25, 0x63, 0xEB)
sp.paragraph_format.space_after = Pt(4)

dp = doc.add_paragraph()
dp.alignment = WD_ALIGN_PARAGRAPH.CENTER
dr = dp.add_run('Collections Hackathon  |  IIIT Campus  |  October 2026')
dr.font.size = Pt(10)
dr.font.color.rgb = RGBColor(0x6B, 0x72, 0x80)
dp.paragraph_format.space_after = Pt(16)

add_divider(doc)

# ═══════════════════════════════════════════════════════════════════════════════
# OVERVIEW: HOW THE SPLIT WORKS
# ═══════════════════════════════════════════════════════════════════════════════

add_heading(doc, 'How the Project is Divided', level=1)

add_para(doc,
    'The complete project spans three hackathon layers and four core innovations. '
    'It is divided into two well-balanced tracks — each self-contained, with a clear handoff point in the middle. '
    'Both members present together in the final 5-minute demo.',
    color='374151')

add_code_block(doc,
    "  MEMBER 1 — DATA & IDENTITY FOUNDATION\n"
    "  =========================================\n"
    "  Layer 1: Data Product Factory\n"
    "       Raw Ingestion -> Silver Curation -> Golden C360\n"
    "  Golden Financial ID  (Entity Resolution + Identity Layer)\n"
    "  Financial Health Engine  (Exposure + Cash-Flow + Behaviour scoring)\n"
    "  Data Contracts + Quality SLA\n"
    "  Governance Framework\n\n"
    "                   [ HANDOFF: Curated C360 + Golden ID + Feature Signals ]\n\n"
    "  MEMBER 2 — INTELLIGENCE & DECISION LAYER\n"
    "  ==========================================\n"
    "  Layer 2: Insight & NLP Platform\n"
    "       NL-to-SQL + Semantic Layer + RAG over Notes/Transcripts\n"
    "  Layer 3: AI Decisioning Engine\n"
    "       Collection Memory + Next Best Action + Smart Payment Plans\n"
    "       Agent Assist Cockpit + Post-Call Summaries + Fair Collections\n"
    "  Presentation & Pitch Deck storytelling")

add_divider(doc)

# ═══════════════════════════════════════════════════════════════════════════════
# MEMBER 1 — FULL DETAIL
# ═══════════════════════════════════════════════════════════════════════════════

add_heading(doc, 'MEMBER 1 — Data & Identity Foundation', level=1, color='1D4ED8')

add_callout(doc,
    'MEMBER 1 SCOPE: ',
    'Layer 1 (Data Product Factory) + Golden Financial ID + Financial Health Engine + '
    'Data Contracts + Governance. This member builds the trusted data foundation that the entire system runs on.',
    bg='EFF6FF', bar='2563EB', label_color='1D4ED8')

# ---- Responsibilities --------------------------------------------------------
add_heading(doc, 'Core Responsibilities', level=2, color='1D4ED8')

add_heading(doc, 'A.  Raw Data Ingestion & Curated Silver Layer', level=3, color='1E3A5F')
add_para(doc, 'Working dataset: All 9 synthetic CSV/JSON sources from the hackathon brief.', color='374151')
for item in [
    'Load all 9 source files: customers, card_accounts, loan_accounts, deposit_accounts, collections_cases, contact_history, agent_notes, call_transcripts, external bureau.',
    'Clean and standardise each source: handle nulls, normalise phone numbers, standardise name formats, fix date columns.',
    'Build a Silver (curated) layer — one clean table per domain, documented with column definitions and data types.',
    'Write a data quality report: null rates, duplicate counts, format violations, freshness.',
]:
    add_bullet(doc, item)

add_heading(doc, 'B.  Entity Resolution & Golden Financial ID', level=3, color='1E3A5F')
add_para(doc, 'The central innovation — linking all source records to one canonical customer identity.', color='374151')
for item in [
    'Implement deterministic matching: exact match on normalised phone, email, or tax-ID fields across the 9 sources.',
    'Implement probabilistic matching: Jaro-Winkler name similarity + address token matching for fuzzy records.',
    'Assign a match confidence score (0.0 – 1.0) to every resolved link.',
    'Generate a Golden Financial ID (GID-XXXXXX) for every resolved customer.',
    'Build the golden_customer_360 master entity table with GID as primary key.',
    'Demonstrate the challenge\'s "One customer, one ID" requirement with metrics: how many source records resolved, average confidence, unmatched records.',
]:
    add_bullet(doc, item)

add_heading(doc, 'C.  Golden Customer 360 Product (Golden C360)', level=3, color='1E3A5F')
for item in [
    'Aggregate key attributes per GID: total exposure, max DPD across portfolio, active facility count, EMI obligations, salary credit trend.',
    'Add lineage metadata to every field: source system, pipeline run timestamp, transformation version.',
    'Mark all PII fields with classification tags (confidential_pii) for access control.',
    'Produce a final golden_customer_360 table ready for Member 2 to query.',
]:
    add_bullet(doc, item)

add_heading(doc, 'D.  Financial Health Engine (Scoring Backend)', level=3, color='1E3A5F')
add_para(doc, 'Computes the 3-branch financial health signal per GID:', color='374151')
for item in [
    'Exposure Score: normalised ratio of total outstanding debt to estimated annual inflow.',
    'Cash-Flow Stress Index: detect missing/late payroll credits, assign LOW / MODERATE / HIGH / SEVERE.',
    'Payment Behaviour Score: PTP fulfilment rate over 24 months, historical DPD velocity.',
    'Composite Financial Health Signal: weighted combination of the three branches.',
    'Write results back to the feature store as computed columns on the Golden C360 table.',
]:
    add_bullet(doc, item)

add_heading(doc, 'E.  Data Contract & Governance Document', level=3, color='1E3A5F')
for item in [
    'Publish one formal Data Contract YAML (Open Data Contract Standard) for the Golden C360 product.',
    'Define quality rules: uniqueness on GID, freshness SLA, positive exposure check, exclusion of protected attributes.',
    'Write a data quality scorecard PDF/slide showing pass/fail status for every rule.',
    'Document AI helpers used: schema drift detection prompt, SQL generation, anomaly detection.',
]:
    add_bullet(doc, item)

# ---- Deliverables --------------------------------------------------------
add_heading(doc, 'Deliverables for Member 1', level=2, color='1D4ED8')

make_table(doc,
    headers=['Deliverable', 'Format', 'Purpose'],
    rows=[
        ('Silver Layer Tables (9 clean domain tables)', 'CSV / SQLite / Parquet', 'Curated source data for Member 2 queries'),
        ('golden_customer_360 master table', 'CSV / Parquet', 'Core data product — primary key = GID'),
        ('Entity Resolution Report', 'Markdown / PDF', 'Metrics: records resolved, confidence distribution'),
        ('Financial Health Feature Table', 'CSV / Parquet', 'Exposure score, cash-flow stress, behaviour score per GID'),
        ('Data Contract YAML', 'YAML file', 'Hackathon mandatory deliverable for Layer 1'),
        ('Data Quality Scorecard', '1 slide / PDF', 'Demonstrates rule checks and SLA compliance'),
        ('Architecture Diagram (Layer 1)', 'Diagram image / slide', 'Included in pitch deck slide 7'),
    ]
)

# ---- Skills & Tools --------------------------------------------------------
add_heading(doc, 'Suggested Tools & Skills', level=2, color='1D4ED8')

make_table(doc,
    headers=['Task', 'Recommended Tool/Library'],
    rows=[
        ('Data loading & cleaning', 'Python + pandas'),
        ('Fuzzy name matching', 'recordlinkage / jellyfish / rapidfuzz'),
        ('SQL transformations', 'SQLite / DuckDB'),
        ('Entity resolution pipeline', 'Python (custom) or splink library'),
        ('Financial health scoring', 'Python: pandas, numpy, scikit-learn (rule-based)'),
        ('Data contract authoring', 'YAML editor / VS Code'),
        ('Visualisations', 'matplotlib / seaborn / plotly'),
    ]
)

add_divider(doc)

# ═══════════════════════════════════════════════════════════════════════════════
# MEMBER 2 — FULL DETAIL
# ═══════════════════════════════════════════════════════════════════════════════

add_heading(doc, 'MEMBER 2 — Intelligence & Decision Layer', level=1, color='059669')

add_callout(doc,
    'MEMBER 2 SCOPE: ',
    'Layer 2 (Insight & NLP) + Layer 3 (AI Decisioning) + Collection Memory + Next Best Action + '
    'Smart Payment Plans + Agent Assist Cockpit + Post-Call Summaries + Pitch Deck storytelling.',
    bg='F0FDF4', bar='059669', label_color='065F46')

# ---- Responsibilities --------------------------------------------------------
add_heading(doc, 'Core Responsibilities', level=2, color='059669')

add_heading(doc, 'A.  Semantic Layer & NL-to-SQL Engine (Layer 2)', level=3, color='1E3A5F')
add_para(doc, 'Consumes the Golden C360 + Silver tables produced by Member 1.', color='374151')
for item in [
    'Build a semantic mapping layer: define human-friendly aliases for all column names (e.g., max_dpd_across_portfolio = "Days Overdue").',
    'Integrate an LLM (Gemini API / OpenAI) to translate plain-English questions into SQL against the C360 schema.',
    'Display the generated SQL transparently alongside the answer.',
    'Implement editable search token UI: user can see and correct misinterpreted filters before execution.',
    'Enforce query guardrails: detect and refuse questions that request protected attributes or out-of-scope data.',
    'Demo benchmark questions: at least 5 working NL queries with results shown end-to-end.',
]:
    add_bullet(doc, item)

add_heading(doc, 'B.  Collection Memory — RAG over Notes & Transcripts (Layer 2)', level=3, color='1E3A5F')
for item in [
    'Chunk and embed agent_notes and call_transcripts using a sentence transformer or Gemini embedding API.',
    'Store vectors in a lightweight vector store (ChromaDB / FAISS / pgvector).',
    'Build a semantic search: given a GID, retrieve the most relevant past interaction context.',
    'Generate a structured Collection Memory summary per customer: last 5 contacts, promises, channel response, outcomes.',
    'Surface the "Five Golden Questions" view: What happened? What is now? What worked? What did not? What next?',
]:
    add_bullet(doc, item)

add_heading(doc, 'C.  Next Best Action Model (Layer 3)', level=3, color='1E3A5F')
for item in [
    'Use the Financial Health signals from Member 1\'s feature table as model input features.',
    'Build or configure a rule-based + ML Next Best Action recommender:',
    '     Channel: SMS / WhatsApp / IVR / Human Call (based on contact history response rates).',
    '     Timing: optimal day-of-week and time-of-day window (from contact_history outcomes).',
    '     Treatment: Reminder / Payment Plan Offer / Hardship Discussion / Specialist Routing.',
    'Ensure every decision is explainable — show the top 2-3 factors that drove the recommendation.',
    'Add a human override mechanism: agent can accept, modify, or reject the NBA suggestion.',
]:
    add_bullet(doc, item)

add_heading(doc, 'D.  Smart Payment Plan Engine (Layer 3)', level=3, color='1E3A5F')
for item in [
    'Consume Debt Service Ratio (DSR) and Financial Health Score from Member 1.',
    'Generate up to 3 tailored payment plan options per delinquent customer:',
    '     Option A: Short-term moratorium (30-day freeze).',
    '     Option B: Step-up amortisation (lower EMIs for 3-6 months, then normal).',
    '     Option C: Conditional fee waiver upon commitment to a 3-instalment plan.',
    'Show plan calculation transparently: exact EMI amounts, duration, total interest impact.',
    'Flag: plans are human-reviewed only — never auto-applied.',
]:
    add_bullet(doc, item)

add_heading(doc, 'E.  Agent Assist Cockpit & Post-Call Summary (Layer 3)', level=3, color='1E3A5F')
for item in [
    'Build a simple Agent Assist UI (web or terminal) showing per customer:',
    '     Golden ID + Identity summary.',
    '     Collection Memory timeline (from Part B).',
    '     Current Financial Health indicator (traffic light: GREEN / AMBER / RED).',
    '     NBA recommendation with explanation.',
    '     Smart Payment Plan options (pending human approval).',
    'Implement automated Post-Call Summary: given a transcript, generate a structured note using an LLM — tagging reason for non-payment, sentiment, promise made, next action.',
    'Optionally add QA scoring: check if agent followed compliance guidelines from transcript.',
]:
    add_bullet(doc, item)

add_heading(doc, 'F.  Pitch Deck & Presentation Storytelling', level=3, color='1E3A5F')
for item in [
    'Own the 10-slide pitch deck — structure and narrative flow.',
    'Integrate architecture diagrams, screenshots, and demo outputs from both members.',
    'Script the 5-minute demo video: clear walkthrough of NL query -> C360 -> Collection Memory -> NBA -> Agent Cockpit.',
    'Prepare the final innovation story: "From Data Silos -> Golden ID -> Fair Decision -> Customer Outcome".',
]:
    add_bullet(doc, item)

# ---- Deliverables --------------------------------------------------------
add_heading(doc, 'Deliverables for Member 2', level=2, color='059669')

make_table(doc,
    headers=['Deliverable', 'Format', 'Purpose'],
    rows=[
        ('Semantic Layer Schema Map', 'JSON / Python dict', 'Alias mapping for NL-to-SQL'),
        ('NL-to-SQL Engine (working demo)', 'Python / Notebook', 'Hackathon Layer 2 core deliverable'),
        ('Collection Memory Vector Store', 'ChromaDB / FAISS index', 'RAG over agent notes & transcripts'),
        ('Next Best Action Recommender', 'Python model / rule engine', 'Hackathon Layer 3 core deliverable'),
        ('Smart Payment Plan Engine', 'Python / Notebook', 'Affordability-driven plan generator'),
        ('Agent Assist Cockpit (demo UI)', 'Streamlit / Gradio / terminal', '5-min demo showcase application'),
        ('Post-Call Summary Generator', 'Python + LLM call', 'Automated structured note from transcript'),
        ('10-Slide Pitch Deck', 'Google Slides / PowerPoint', 'Hackathon submission'),
        ('5-Minute Demo Video Script', 'Document', 'Guides video recording'),
    ]
)

# ---- Skills & Tools --------------------------------------------------------
add_heading(doc, 'Suggested Tools & Skills', level=2, color='059669')

make_table(doc,
    headers=['Task', 'Recommended Tool/Library'],
    rows=[
        ('LLM integration (NL-to-SQL, summaries)', 'Gemini API / OpenAI API / LangChain'),
        ('Vector embeddings', 'Gemini Embedding / sentence-transformers'),
        ('Vector store', 'ChromaDB / FAISS'),
        ('NBA model', 'scikit-learn (GBM) or rule engine (Python dict)'),
        ('Agent UI', 'Streamlit / Gradio'),
        ('SQL query execution', 'DuckDB / SQLite (shared with Member 1)'),
        ('Presentation', 'Google Slides / Canva / PowerPoint'),
    ]
)

add_divider(doc)

# ═══════════════════════════════════════════════════════════════════════════════
# HANDOFF PROTOCOL
# ═══════════════════════════════════════════════════════════════════════════════

add_heading(doc, 'Handoff Protocol Between Members', level=1)

add_para(doc,
    'Member 1 produces the data foundation. Member 2 consumes it. '
    'The handoff is a shared set of files/tables in an agreed folder structure.',
    color='374151')

add_code_block(doc,
    "  project/\n"
    "  |-- data/\n"
    "  |   |-- raw/           <- Member 1: original CSVs\n"
    "  |   |-- silver/        <- Member 1: cleaned per-domain tables\n"
    "  |   |-- golden/        <- Member 1: golden_customer_360.csv + features\n"
    "  |\n"
    "  |-- contracts/\n"
    "  |   |-- collections_360_contract.yaml   <- Member 1\n"
    "  |\n"
    "  |-- nlp/               <- Member 2: NL-to-SQL engine\n"
    "  |-- memory/            <- Member 2: Collection Memory RAG\n"
    "  |-- decisioning/       <- Member 2: NBA + Payment Plans\n"
    "  |-- ui/                <- Member 2: Agent Assist Cockpit\n"
    "  |\n"
    "  |-- docs/\n"
    "      |-- pitch_deck.pptx\n"
    "      |-- architecture.docx")

add_callout(doc,
    'CRITICAL HANDOFF FILES: ',
    'Member 1 must produce golden_customer_360.csv and financial_health_features.csv '
    'before Member 2 can build the NBA model and Smart Payment Plans. '
    'Coordinate on schema early so Member 2 can start on NL-to-SQL in parallel.',
    bg='FEF3C7', bar='F59E0B', label_color='92400E')

add_divider(doc)

# ═══════════════════════════════════════════════════════════════════════════════
# PARALLEL WORK SCHEDULE
# ═══════════════════════════════════════════════════════════════════════════════

add_heading(doc, 'Suggested Parallel Work Schedule', level=1)

add_para(doc,
    'Both members can begin simultaneously. Member 1 sets up the data pipeline while '
    'Member 2 sets up the LLM and NL engine against a minimal sample dataset.',
    color='374151')

make_table(doc,
    headers=['Phase', 'Member 1 Activity', 'Member 2 Activity'],
    rows=[
        ('Phase 1\n(Start)', 'Load all 9 raw CSVs; inspect schemas; plan Silver tables', 'Set up LLM API; draft 5 benchmark NL questions; build semantic alias map on sample data'),
        ('Phase 2', 'Build Silver layer cleaning scripts; run entity resolution on full dataset; generate GIDs', 'Implement NL-to-SQL engine; test against sample C360 schema; set up ChromaDB for RAG'),
        ('Phase 3\n(Handoff)', 'Produce final golden_customer_360.csv + financial_health_features.csv; write Data Contract', 'Embed agent_notes + transcripts; build Collection Memory retrieval; test Five Golden Questions'),
        ('Phase 4', 'Write Data Quality Scorecard; finalise architecture diagram for pitch deck', 'Build NBA recommender + Smart Payment Plan engine using Member 1\'s feature table'),
        ('Phase 5', 'Assist with governance slide; review Member 2\'s demo for data accuracy', 'Build Agent Assist Cockpit UI; record 5-min demo video; finalise 10-slide pitch deck'),
    ]
)

add_divider(doc)

# ═══════════════════════════════════════════════════════════════════════════════
# HACKATHON JUDGING ALIGNMENT
# ═══════════════════════════════════════════════════════════════════════════════

add_heading(doc, 'Judging Criteria Ownership', level=1)

add_para(doc,
    'The hackathon scores 5 criteria. Here is which member leads each dimension:',
    color='374151')

make_table(doc,
    headers=['Judging Criterion', 'Weight', 'Lead Owner', 'Supporting Member'],
    rows=[
        ('Business Understanding', '20%', 'Member 2 (Pitch Deck + Storytelling)', 'Member 1 (Data context)'),
        ('Technical Quality & Working Demo', '30%', 'Both equally', 'Member 1: Layer 1 pipeline; Member 2: Layer 2+3 demo'),
        ('Innovative Ideas', '15%', 'Both equally', 'Golden ID (M1) + Collection Memory + Fair Collections (M2)'),
        ('Teamwork', '15%', 'Both — visible collaboration', 'Handoff quality and code integration'),
        ('Presentation & Storytelling', '20%', 'Member 2 (Deck + Demo script)', 'Member 1 (Data architecture visual)'),
    ]
)

add_divider(doc)

# ═══════════════════════════════════════════════════════════════════════════════
# SHARED RESPONSIBILITIES
# ═══════════════════════════════════════════════════════════════════════════════

add_heading(doc, 'Shared Responsibilities (Both Members)', level=1)

for item in [
    'Code repository setup: GitHub repo with clear folder structure from day one.',
    'README.md with project overview, instructions, and demo run guide.',
    'Final architecture diagram review together before submission.',
    'Code repository submission: ensure both member contributions are committed.',
    'Final 5-minute demo walkthrough rehearsal together.',
    'Mutual review of each other\'s output before final submission.',
]:
    add_bullet(doc, item)

add_divider(doc)

# ═══════════════════════════════════════════════════════════════════════════════
# SUMMARY CARD
# ═══════════════════════════════════════════════════════════════════════════════

add_heading(doc, 'Quick Reference: Who Does What', level=1)

add_code_block(doc,
    "  MEMBER 1 — DATA & IDENTITY FOUNDATION\n"
    "  ---------------------------------------\n"
    "  [ Layer 1 ]  Raw Ingestion -> Silver -> Golden C360\n"
    "  [ Identity ] Entity Resolution -> Golden Financial ID\n"
    "  [ Health  ]  Financial Health Engine (Exposure + Cash-Flow + Behaviour)\n"
    "  [ Govern  ]  Data Contract YAML + Quality Scorecard\n\n"
    "  MEMBER 2 — INTELLIGENCE & DECISION LAYER\n"
    "  ------------------------------------------\n"
    "  [ Layer 2 ]  Semantic Layer + NL-to-SQL + RAG (Collection Memory)\n"
    "  [ Layer 3 ]  NBA + Smart Payment Plans + Agent Assist Cockpit\n"
    "  [ UX      ]  Post-Call Summaries + QA Automation\n"
    "  [ Pres.   ]  10-Slide Pitch Deck + 5-Min Demo\n\n"
    "  SHARED\n"
    "  -------\n"
    "  GitHub Repo + README + Architecture Review + Final Demo Rehearsal")

add_para(doc,
    'Both members present the final solution together. The combined output delivers '
    'all three hackathon layers, the Golden Financial ID innovation, Collection Memory, '
    'Fair Collections Intelligence, and a working end-to-end demo.',
    color='374151', space_before=8)

# ═══════════════════════════════════════════════════════════════════════════════
# SAVE
# ═══════════════════════════════════════════════════════════════════════════════

output_path = r'C:\Users\dadid\.gemini\antigravity\scratch\collections-ai\Team_Division_Plan.docx'
doc.save(output_path)
print(f'Saved -> {output_path}')
