from docx import Document
from docx.shared import Pt, RGBColor, Inches, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
import copy

# ── helpers ───────────────────────────────────────────────────────────────────

def shade_cell(cell, fill_hex):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), fill_hex)
    tcPr.append(shd)

def set_cell_border(cell, **kwargs):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcBorders = OxmlElement('w:tcBorders')
    for side in ['top', 'left', 'bottom', 'right']:
        border = OxmlElement(f'w:{side}')
        border.set(qn('w:val'), kwargs.get('val', 'single'))
        border.set(qn('w:sz'), kwargs.get('sz', '4'))
        border.set(qn('w:space'), '0')
        border.set(qn('w:color'), kwargs.get('color', '2563EB'))
        tcBorders.append(border)
    tcPr.append(tcBorders)

def add_heading(doc, text, level=1, color='1E3A5F'):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = p.add_run(text)
    run.bold = True
    if level == 1:
        run.font.size = Pt(18)
    elif level == 2:
        run.font.size = Pt(14)
    elif level == 3:
        run.font.size = Pt(12)
    r, g, b = tuple(int(color[i:i+2], 16) for i in (0, 2, 4))
    run.font.color.rgb = RGBColor(r, g, b)
    p.paragraph_format.space_before = Pt(14)
    p.paragraph_format.space_after = Pt(4)
    return p

def add_para(doc, text, bold=False, italic=False, size=10.5, color='222222',
             align=WD_ALIGN_PARAGRAPH.LEFT, space_before=2, space_after=4):
    p = doc.add_paragraph()
    p.alignment = align
    run = p.add_run(text)
    run.bold = bold
    run.italic = italic
    run.font.size = Pt(size)
    r, g, b = tuple(int(color[i:i+2], 16) for i in (0, 2, 4))
    run.font.color.rgb = RGBColor(r, g, b)
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.space_after = Pt(space_after)
    return p

def add_bullet(doc, text, level=0, color='1E3A5F'):
    p = doc.add_paragraph(style='List Bullet')
    run = p.add_run(text)
    run.font.size = Pt(10)
    r, g, b = tuple(int(color[i:i+2], 16) for i in (0, 2, 4))
    run.font.color.rgb = RGBColor(r, g, b)
    p.paragraph_format.left_indent = Inches(0.2 + level * 0.25)
    p.paragraph_format.space_after = Pt(2)
    return p

def add_code_block(doc, text):
    """Monospaced architecture/flow block in a shaded box."""
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.3)
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(4)
    run = p.add_run(text)
    run.font.name = 'Courier New'
    run.font.size = Pt(8.5)
    run.font.color.rgb = RGBColor(0x1E, 0x3A, 0x5F)
    # light blue-grey shade
    pPr = p._p.get_or_add_pPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), 'EEF2FF')
    pPr.append(shd)
    return p

def add_divider(doc):
    p = doc.add_paragraph()
    pPr = p._p.get_or_add_pPr()
    pBdr = OxmlElement('w:pBdr')
    bottom = OxmlElement('w:bottom')
    bottom.set(qn('w:val'), 'single')
    bottom.set(qn('w:sz'), '4')
    bottom.set(qn('w:space'), '1')
    bottom.set(qn('w:color'), '2563EB')
    pBdr.append(bottom)
    pPr.append(pBdr)
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(6)

def add_callout(doc, text, bg='FFF3CD', bar='F59E0B', label='NOTE'):
    """Callout box: coloured left-bar table row."""
    table = doc.add_table(rows=1, cols=2)
    table.columns[0].width = Cm(0.45)
    table.columns[1].width = Cm(15.55)
    lc = table.cell(0, 0)
    rc = table.cell(0, 1)
    shade_cell(lc, bar)
    shade_cell(rc, bg)
    lc._tc.get_or_add_tcPr()
    p = rc.paragraphs[0]
    run = p.add_run(f'{label}  ')
    run.bold = True
    run.font.size = Pt(9)
    run.font.color.rgb = RGBColor(0x92, 0x40, 0x00)
    run2 = p.add_run(text)
    run2.font.size = Pt(9.5)
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.left_indent = Inches(0.08)
    doc.add_paragraph().paragraph_format.space_after = Pt(4)

# ── document setup ─────────────────────────────────────────────────────────────

doc = Document()

# Page margins
section = doc.sections[0]
section.top_margin    = Cm(2.0)
section.bottom_margin = Cm(2.0)
section.left_margin   = Cm(2.5)
section.right_margin  = Cm(2.0)

# Default font
style = doc.styles['Normal']
style.font.name = 'Calibri'
style.font.size = Pt(10.5)

# ═══════════════════════════════════════════════════════════════════════════════
# COVER / TITLE BLOCK
# ═══════════════════════════════════════════════════════════════════════════════

tp = doc.add_paragraph()
tp.alignment = WD_ALIGN_PARAGRAPH.CENTER
tr = tp.add_run('Enterprise Collections Intelligence Platform')
tr.bold = True
tr.font.size = Pt(24)
tr.font.color.rgb = RGBColor(0x1E, 0x3A, 0x5F)
tp.paragraph_format.space_before = Pt(10)
tp.paragraph_format.space_after  = Pt(6)

sp = doc.add_paragraph()
sp.alignment = WD_ALIGN_PARAGRAPH.CENTER
sr = sp.add_run('Unified Architecture · Golden Financial Identity · Fair Collections Framework')
sr.italic = True
sr.font.size = Pt(13)
sr.font.color.rgb = RGBColor(0x25, 0x63, 0xEB)
sp.paragraph_format.space_after = Pt(16)

add_divider(doc)

# ═══════════════════════════════════════════════════════════════════════════════
# SECTION 0 — EXECUTIVE OVERVIEW
# ═══════════════════════════════════════════════════════════════════════════════

add_heading(doc, '0.  Executive Overview & The Innovation Paradigm', level=1)

add_para(doc,
    'Modern debt collection is fundamentally broken by data fragmentation and operational myopia. '
    'Conventional banking collections typically ask only one narrow, punitive question:')

add_para(doc,
    '"How much does this customer owe on this specific overdue account, and how fast can we extract payment?"',
    italic=True, color='B45309')

add_para(doc,
    'Our solution transforms that question entirely into:')

add_para(doc,
    '"Who is this customer financially, what has happened before, what is their current relevant '
    'financial context, and what is the most appropriate, sustainable next interaction?"',
    bold=True, italic=True, color='1E3A5F')

add_para(doc,
    'This is the foundational difference between Collections Automation (robotic dialers, rigid dunning) '
    'and Customer-Centric Collections Intelligence.')

add_code_block(doc,
    "TRADITIONAL MODEL:\n"
    "  Data Silos → Overdue Loan Record → Rigid Dunning → Aggressive Recovery / Default\n\n"
    "OUR INNOVATIVE ARCHITECTURE:\n"
    "  Golden Financial ID\n"
    "          ↓\n"
    "  Customer 360  (Governed Data Products)\n"
    "          ↓\n"
    "  Collection Memory  (Cross-Cycle Context)\n"
    "          ↓\n"
    "  Financial Health Engine  (Exposure + Cash Flow + Behaviour)\n"
    "          ↓\n"
    "  Fair Decision Intelligence  (NBA + Smart Payment Plans)\n"
    "          ↓\n"
    "  Human-in-the-Loop Review\n"
    "          ↓\n"
    "  Appropriate Customer Interaction & Sustainable Outcome\n"
    "          ↓\n"
    "  Continuous Feedback Loop  ↺")

add_divider(doc)

# ═══════════════════════════════════════════════════════════════════════════════
# SECTION 1 — GOLDEN FINANCIAL IDENTITY LAYER
# ═══════════════════════════════════════════════════════════════════════════════

add_heading(doc, '1.  Core Innovation: The Golden Financial Identity Layer', level=1)

# 1.1
add_heading(doc, '1.1  Conceptual Definition & Responsible Analogy', level=2)

add_para(doc,
    'Inspired by the persistent identity principle of government-issued identifiers (such as PAN), '
    'we introduce the Golden Financial ID — a conceptual financial data identity layer.')

add_para(doc,
    'It enables participating financial institutions and authorised ecosystem systems to associate all '
    'relevant, consented, and permitted financial relationships belonging to the same individual into a '
    'unified identity fabric.', color='374151')

add_callout(doc,
    'HACKATHON PROTOTYPE: Implemented as an entity-resolved master key across the synthetic challenge datasets '
    '(cards, loans, deposits, collections, contact history, notes).\n'
    'FUTURE ENTERPRISE VISION: A privacy-preserving, consent-mediated, interoperable financial identity layer '
    'across authorised financial ecosystem participants and Account Aggregators.\n'
    'EXPLICIT DISCLAIMER: We DO NOT claim that our prototype creates a real-world nationwide identity '
    'infrastructure. It is a strictly governed internal software architecture.',
    bg='EFF6FF', bar='2563EB', label='GOVERNANCE BOUNDARY')

# 1.2
add_heading(doc, '1.2  The Fragmentation Dilemma', level=2)

add_para(doc,
    'In enterprise banking, the same individual exists as disconnected, inconsistent records:',
    color='374151')

add_code_block(doc,
    "  Loan System Record\n"
    "      ≠\n"
    "  Card System Record\n"
    "      ≠\n"
    "  Deposit System Record\n"
    "      ≠\n"
    "  CRM Customer Profile\n\n"
    "  (Even when they represent the SAME living individual)")

add_para(doc,
    'The Golden Financial ID resolves this through an Automated Entity Resolution Pipeline:',
    color='374151')

add_code_block(doc,
    " ┌──────────────────┐   ┌──────────────────┐   ┌──────────────────┐\n"
    " │ Card Sys Record  │   │ Loan Sys Record  │   │ Deposit Acct Rec │\n"
    " └────────┬─────────┘   └────────┬─────────┘   └────────┬─────────┘\n"
    "          │                      │                      │\n"
    "          ▼                      ▼                      ▼\n"
    " ┌──────────────────────────────────────────────────────────────────┐\n"
    " │           HYBRID ENTITY RESOLUTION PIPELINE                      │\n"
    " │  • Deterministic Matching  (Tax ID, Gov ID, Normalised Phone)    │\n"
    " │  • Probabilistic Matching  (Jaro-Winkler, Address Parsing)       │\n"
    " │  • Graph-Based Household & Relationship Linkage                  │\n"
    " │  • Match Confidence Scoring  (≥ 0.92 = Auto Golden Link)         │\n"
    " └───────────────────────────────┬──────────────────────────────────┘\n"
    "                                 │\n"
    "                                 ▼\n"
    " ┌──────────────────────────────────────────────────────────────────┐\n"
    " │              GOLDEN FINANCIAL ID  (e.g., GID-849204)             │\n"
    " │          The Unified Identity Anchor for All Data Products       │\n"
    " └───────────────────────────────┬──────────────────────────────────┘\n"
    "                                 │\n"
    "                                 ▼\n"
    " ┌──────────────────────────────────────────────────────────────────┐\n"
    " │                   UNIFIED FINANCIAL CONTEXT                      │\n"
    " └──────────────────────────────────────────────────────────────────┘")

add_para(doc,
    'This elevates the hackathon\'s "One customer, one ID" requirement from a simple database join '
    'into an enterprise-grade, reusable financial identity layer.', bold=False, color='374151')

# 1.3
add_heading(doc, '1.3  Golden ID as the Root of Customer 360 (Not a Monolithic Database)', level=2)

add_para(doc,
    'The Golden Financial ID does NOT contain every piece of customer data. It functions as a trusted '
    'identity key that connects governed, distributed data products.', bold=True, color='1E3A5F')

add_code_block(doc,
    "  Golden Financial ID  (Identity Anchor / Trusted Key)\n"
    "          ↓\n"
    "  Collections 360  (Curated Core Domain)\n"
    "          ↓\n"
    "  Financial Relationship  (Loans, Cards, Deposits)\n"
    "          ↓\n"
    "  Interaction History  (Omnichannel Contacts & Outcomes)\n"
    "          ↓\n"
    "  Financial Health  (Affordability, Exposure, Cash-Flow Stress)\n"
    "          ↓\n"
    "  Decision Intelligence  (Next Best Action & Fair Guidance)")

add_para(doc, 'Why this separation of concerns is mission-critical:', bold=True, color='1E3A5F')
for pt in [
    'Preserves Data Ownership — each source system retains ownership of its domain data.',
    'Access Control (RBAC/ABAC) — different agent personas query only authorised data products.',
    'End-to-End Lineage — every field retains its source origin and transformation timestamp.',
    'Data Privacy & Compliance — PII attack surface is minimised; unconsented attributes are never joined.',
    'Modularity & Scalability — source systems upgrade independently without breaking downstream consumers.',
]:
    add_bullet(doc, pt)

# 1.4
add_heading(doc, '1.4  Golden Financial ID: Conceptual Data View', level=2)

add_callout(doc,
    'Attributes are resolved and displayed ONLY where available, relevant, legally consented/permitted, '
    'and strictly governed.',
    bg='F0FDF4', bar='16A34A', label='ARCHITECTURAL PRINCIPLE')

add_code_block(doc,
    "GOLDEN FINANCIAL ID  (e.g., GID-849204)\n"
    "│\n"
    "├── Identity / Customer Profile  (Governed PII Domain)\n"
    "│   ├── Standardised Full Name & Demographics\n"
    "│   ├── Verified Contact Coordinates  (Phone, Email, Address)\n"
    "│   ├── Consent & Communication Preferences  (DND, preferred channel, timezone)\n"
    "│   └── Lineage Metadata  (Match confidence: 0.98, last verified timestamp)\n"
    "│\n"
    "├── Employment & Income Signals  (Cash-Flow Domain)\n"
    "│   ├── Latest verified monthly income\n"
    "│   ├── Salary credit consistency trends  (3-month variance, deposit frequency)\n"
    "│   ├── Employment verification status\n"
    "│   └── Income disruption flags  (absence of expected payroll credit)\n"
    "│\n"
    "├── Loan Facilities  (Lending Domain)\n"
    "│   ├── Active term loans  (Home, Auto, Personal, Education)\n"
    "│   ├── Total outstanding principal balance\n"
    "│   ├── Aggregated monthly EMI obligations\n"
    "│   └── Delinquency status  (DPD bucket, overdue balance, roll-rate history)\n"
    "│\n"
    "├── Revolving Credit Cards  (Card Domain)\n"
    "│   ├── Active card accounts & credit limits\n"
    "│   ├── Current aggregate outstanding balance\n"
    "│   ├── Credit utilisation %  (e.g., 88% = high stress)\n"
    "│   └── Min Due vs Full Statement payment behaviour history\n"
    "│\n"
    "├── Deposit Accounts & Liquidity  (Treasury / Deposit Domain)\n"
    "│   ├── Liquid deposit balance trendlines\n"
    "│   ├── Inflow vs outflow ratios  (net liquidity buffer)\n"
    "│   └── Recurring debit obligations  (utility mandates, standing instructions)\n"
    "│\n"
    "├── Asset & Secured Property Signals  (Future Capability — Collateral Domain)\n"
    "│   └── Where legally recorded, verified, and permitted by regulatory framework\n"
    "│\n"
    "├── Collections History  (Case Management Domain)\n"
    "│   ├── Historical case lifecycles  (resolved, restructured, settled)\n"
    "│   ├── Promises-to-Pay history  (count promised, fulfilled vs broken)\n"
    "│   └── Previous strategy assignments, hardship concessions, recovery outcomes\n"
    "│\n"
    "└── Interaction Memory & Unstructured Signals  (Interaction Domain)\n"
    "    ├── Cross-channel contact timeline  (telephony, SMS, WhatsApp, mobile push)\n"
    "    ├── Agent qualitative observation notes\n"
    "    ├── Summarised call transcripts  (sentiment, dispute flags, reason for non-payment)\n"
    "    └── Speech emotion & customer distress indicators")

# 1.5
add_heading(doc, '1.5  Incremental Customer Update Layer (Living 360, Not Stale Snapshot)', level=2)

add_para(doc,
    'Traditional systems operate on static, month-end batch snapshots. Under our architecture, the '
    'Golden Financial ID anchors an Incremental Customer Update Layer:', color='374151')

add_code_block(doc,
    "  OLD CUSTOMER SNAPSHOT\n"
    "          +\n"
    "  NEW VERIFIED DATA  (salary credit, UPI payment, new loan, dispute logged)\n"
    "          ↓\n"
    "  INCREMENTAL RECONCILIATION ENGINE\n"
    "          ↓\n"
    "  UPDATED LIVING CUSTOMER CONTEXT\n"
    "          ↓\n"
    "  REAL-TIME NEXT BEST ACTION RE-SCORING")

for pt in [
    'Income & Payroll: A detected payroll delay immediately revises the cash-flow stress score, preventing harsh calls.',
    'Intra-Day Payments: When a UPI payment clears, all reminder queues for that Golden ID halt within seconds.',
    'Interaction Delta: A note logged at 10:00 AM updates context for the outbound dialer scheduled at 2:00 PM.',
]:
    add_bullet(doc, pt)

add_divider(doc)

# ═══════════════════════════════════════════════════════════════════════════════
# SECTION 2 — COLLECTION MEMORY
# ═══════════════════════════════════════════════════════════════════════════════

add_heading(doc, '2.  Core Innovation: Collection Memory & The Agent Experience', level=1)

add_heading(doc, '2.1  Philosophy: "Never Make the Next Agent Start from Zero"', level=2)

add_para(doc,
    'Collection Memory is a persistent, semantic, structured memory repository that binds every '
    'interaction — across time, channels, and agents — to the customer\'s Golden Financial ID. '
    'It comprehensively captures:', color='374151')

for item in [
    'Previous contact attempts (channel, date, exact timestamps)',
    'Communication channel efficacy (SMS, WhatsApp, telephony, mobile push)',
    'Customer pick-up rates and response patterns',
    'Promises-to-Pay (PTP): promised dates, amounts, fulfilment vs breach',
    'Historical treatment strategies and hardship concessions',
    'Qualitative agent notes and unstructured observations',
    'Synthesised conversational transcripts and sentiment signals',
    'AI recommendations offered vs. human agent overrides',
    'Final case lifecycle outcomes',
]:
    add_bullet(doc, item)

# 2.2
add_heading(doc, '2.2  The Collection Memory Agent Application Flow', level=2)

add_code_block(doc,
    "  CUSTOMER QUEUE\n"
    "          ↓\n"
    "  GOLDEN FINANCIAL ID\n"
    "          ↓\n"
    "  COLLECTION MEMORY\n"
    "          ↓\n"
    "  CURRENT FINANCIAL CONTEXT\n"
    "          ↓\n"
    "  FINANCIAL HEALTH\n"
    "          ↓\n"
    "  NEXT BEST ACTION")

add_para(doc, 'The Five Golden Questions — answered instantly for every agent:', bold=True, color='1E3A5F')

# Table for 5 golden questions
tbl = doc.add_table(rows=6, cols=2)
tbl.style = 'Table Grid'
headers = ['Key Agent Question', 'Information Surfaced by the Platform']
for i, h in enumerate(headers):
    cell = tbl.cell(0, i)
    shade_cell(cell, '1E3A5F')
    run = cell.paragraphs[0].add_run(h)
    run.bold = True
    run.font.size = Pt(10)
    run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)

rows = [
    ('"What happened before?"',
     'Contact timeline of last 5 attempts; last promise was ₹12,000 on 15th (broken — delayed medical reimbursement noted in transcript).'),
    ('"What is happening now?"',
     'Current DPD = 42; 2 open obligations (Card: ₹45,000 overdue; Personal Loan: regular); salary credited 4 days late this month.'),
    ('"What has worked?"',
     'Customer responds promptly to WhatsApp 6 – 8 PM; courteous phone discussions yielded 80% PTP fulfilment in prior cycles.'),
    ('"What has NOT worked?"',
     'Automated morning IVR calls went unanswered 4 times; demanding full bullet payment caused call abandonment.'),
    ('"What should I do next?"',
     'NBA: Propose 3-month split instalment plan; waive late fee upon first instalment; avoid harsh legal dunning tone.'),
]
for i, (q, a) in enumerate(rows, start=1):
    shade_cell(tbl.cell(i, 0), 'EFF6FF')
    shade_cell(tbl.cell(i, 1), 'FAFAFA')
    qr = tbl.cell(i, 0).paragraphs[0].add_run(q)
    qr.bold = True
    qr.font.size = Pt(9.5)
    qr.font.color.rgb = RGBColor(0x1E, 0x3A, 0x5F)
    ar = tbl.cell(i, 1).paragraphs[0].add_run(a)
    ar.font.size = Pt(9.5)

doc.add_paragraph().paragraph_format.space_after = Pt(6)
add_divider(doc)

# ═══════════════════════════════════════════════════════════════════════════════
# SECTION 3 — FINANCIAL HEALTH ENGINE
# ═══════════════════════════════════════════════════════════════════════════════

add_heading(doc, '3.  Core Innovation: Financial Health Engine & Exposure Analysis', level=1)

add_heading(doc, '3.1  Beyond One-Dimensional DPD Buckets', level=2)

add_para(doc,
    'Traditional collections categorise customers purely by Days Past Due (DPD) bucket, treating a '
    'tech professional who forgot a card payment identically to a worker experiencing a catastrophic '
    'wage cut. Our Financial Health Engine answers:', color='374151')

add_para(doc,
    '"Can the customer\'s current financial situation reasonably support the existing payment obligation?"',
    bold=True, italic=True, color='1E3A5F')

add_callout(doc,
    'Financial Health is an AI-assisted decision-support SIGNAL, not a definitive legal or medical '
    'diagnosis. It empowers human judgement with transparent evidence.',
    bg='FEF3C7', bar='F59E0B', label='GOVERNANCE NOTICE')

# Three-branch analysis diagram
add_heading(doc, '3.2  The Three-Branch Financial Analysis Architecture', level=2)

add_code_block(doc,
    "                            CUSTOMER DATA\n"
    "                                  ↓\n"
    "                           COLLECTIONS 360\n"
    "                                  ↓\n"
    "                          FINANCIAL ANALYSIS\n"
    "                                  ↓\n"
    "         ┌────────────────────────┼────────────────────────┐\n"
    "         ↓                        ↓                        ↓\n"
    "  EXPOSURE ANALYSIS        CASH-FLOW STRESS         PAYMENT BEHAVIOUR\n"
    "  (Multi-Product Debt)    (Payroll Disruption)     (Promise Track Record)\n"
    "         └────────────────────────┼────────────────────────┘\n"
    "                                  ↓\n"
    "                          FINANCIAL HEALTH\n"
    "                                  ↓\n"
    "                       DECISION INTELLIGENCE\n"
    "                                  ↓\n"
    "                          NEXT BEST ACTION\n"
    "                                  ↓\n"
    "         ┌────────────────────────┼────────────────────────┐\n"
    "         ↓                        ↓                        ↓\n"
    "   REMINDER-FIRST          SMART PAYMENT PLAN       SPECIALIST\n"
    "     (Digital Nudge)       (Affordability-Driven)   HUMAN REVIEW")

add_heading(doc, 'Branch 1: Exposure Analysis (Total Obligation Mapping)', level=3)
for pt in [
    'Aggregate outstanding debt across all credit lines, unsecured loans, and credit cards.',
    'Cumulative monthly debt-service burden — sum of all active EMIs.',
    'Credit card utilisation: whether lines are maxed (>90%), indicating reliance on revolving credit for living expenses.',
    'Ecosystem exposure: whether the delinquent loan represents 5% or 80% of the customer\'s total obligations.',
]:
    add_bullet(doc, pt)

add_heading(doc, 'Branch 2: Cash-Flow Stress (Liquidity Shock Detection)', level=3)
for pt in [
    'Income Trajectory: 6-month trailing average of net pay deposits.',
    'Salary Disruption: detecting delayed payroll (normally credited on 30th, absent by 7th).',
    'Buffer Ratio: average end-of-month liquid balance vs. upcoming EMI requirements.',
]:
    add_bullet(doc, pt)

add_para(doc, 'Real-World Contrast:', bold=True, color='1E3A5F', space_before=6)

add_code_block(doc,
    " ┌────────────────────────────────┬────────────────────────────────┐\n"
    " │  CUSTOMER A: Admin Oversight   │  CUSTOMER B: Acute Shock       │\n"
    " ├────────────────────────────────┼────────────────────────────────┤\n"
    " │  Overdue: ₹25,000 (35 DPD)    │  Overdue: ₹25,000 (35 DPD)    │\n"
    " │  Income:  ₹1,80,000/mo stable │  Income:  Drops from ₹80k → 0 │\n"
    " │  Cash:    Consistent credit   │  Cash:    Payroll missing 45d  │\n"
    " │  Exposure: Low (<15% income)  │  Exposure: High (3 active EMIs)│\n"
    " │  Behaviour: Always pays full  │  Behaviour: Consistent → halt  │\n"
    " ├────────────────────────────────┼────────────────────────────────┤\n"
    " │  → Digital Reminder Nudge     │  → Proactive Hardship Plan /   │\n"
    " │                               │    Specialist Human Review     │\n"
    " └────────────────────────────────┴────────────────────────────────┘")

add_heading(doc, 'Branch 3: Payment Behaviour (Willingness & Reliability Profiling)', level=3)
for pt in [
    'Historical Delinquency Velocity: how often the customer rolls from 30 DPD to 60 DPD.',
    'Promise-to-Pay (PTP) Integrity: ratio of honoured promises vs. broken over 24 months.',
    'Channel Affinity: probability of engagement via WhatsApp, IVR, Email, or Direct Call.',
]:
    add_bullet(doc, pt)

add_heading(doc, '3.3  Income-to-Obligation Affordability Architecture', level=2)

add_para(doc,
    'Instead of demanding full settlement regardless of viability, the system models affordability dynamically. '
    'No rigid universal thresholds are imposed. An institution-configurable policy engine allows risk and '
    'credit policy teams to define affordability rules by segment, geography, and product type.',
    color='374151')

add_para(doc,
    'Estimated Debt Service Ratio (DSR) = Total Monthly EMIs + Minimum Card Dues  ÷  Verified Monthly Inflow',
    bold=True, color='2563EB')

add_divider(doc)

# ═══════════════════════════════════════════════════════════════════════════════
# SECTION 4 — FAIR COLLECTIONS & SMART PAYMENT PLANS
# ═══════════════════════════════════════════════════════════════════════════════

add_heading(doc, '4.  Core Innovation: Fair Collections Intelligence & Smart Payment Plans', level=1)

add_heading(doc, '4.1  The Core Operating Principle', level=2)

add_callout(doc,
    '"Collections intelligence should optimize not only for recovery, but also for the appropriateness '
    'and sustainability of the customer interaction."',
    bg='EFF6FF', bar='2563EB', label='NORTH STAR PRINCIPLE')

add_code_block(doc,
    "  GOVERNED DATA  +  FINANCIAL HEALTH  +  INTERACTION MEMORY\n"
    "                              ↓\n"
    "              RESPONSIBLE AI DECISION ENGINE\n"
    "                              ↓\n"
    "                  EMPOWERED HUMAN JUDGEMENT\n"
    "                              ↓\n"
    "           ETHICAL, SUSTAINABLE CUSTOMER OUTCOME")

add_heading(doc, '4.2  Smart Payment Plan Recommendation Flow', level=2)

add_code_block(doc,
    "  Income / Cash-Flow\n"
    "          +\n"
    "  Existing EMI Obligations\n"
    "          +\n"
    "  Outstanding Exposure\n"
    "          +\n"
    "  Payment Behaviour\n"
    "          +\n"
    "  Current Delinquency\n"
    "          ↓\n"
    "  Affordability & Financial Health Signals\n"
    "          ↓\n"
    "  Explainable Payment Plan Options\n"
    "  (Short-Term Freeze | Step-Up Instalments | Conditional Fee Waiver)\n"
    "          ↓\n"
    "  Mandatory Human Agent Review & Authorisation\n"
    "          ↓\n"
    "  Empathetic Agent-Customer Discussion\n"
    "          ↓\n"
    "  Sustainable Commitment & Incremental System Update")

add_para(doc, 'Ground Rules:', bold=True, color='1E3A5F')
for pt in [
    'Never Automatically Imposed — AI generates options; only an authorised human agent can execute them.',
    'Explainable Terms — agent can view the exact formula showing why the proposed instalment fits cash flow.',
    'Hardship Flagging — medical distress or bereavement detected in transcripts diverts cases to Vulnerability Teams, away from robotic dialers.',
]:
    add_bullet(doc, pt)

add_heading(doc, '4.3  Future Horizon: Asset-Backed Restructuring', level=2)

add_callout(doc,
    'The prototype does NOT seize, auto-liquidate, or unilaterally hypothecate customer assets. '
    'Where collateral is legitimately registered and legally eligible, authorised restructuring officers '
    'may use future platform capabilities to evaluate secured refinancing. All decisions remain with '
    'licensed credit professionals.',
    bg='FEF3C7', bar='F59E0B', label='FUTURE CONCEPT — STRICTLY GOVERNED')

add_heading(doc, '4.4  Future Horizon: Collections-to-Advisory Transition', level=2)

add_code_block(doc,
    "  Collections Case Resolution\n"
    "          ↓\n"
    "  Financial Health Assessment\n"
    "          ↓\n"
    "  Proactive Financial Advisory & Credit Repair Guidance")

add_divider(doc)

# ═══════════════════════════════════════════════════════════════════════════════
# SECTION 5 — MASTER ARCHITECTURE
# ═══════════════════════════════════════════════════════════════════════════════

add_heading(doc, '5.  Master End-to-End Architecture Diagram', level=1)

add_code_block(doc,
    "                          BANK DATA SOURCES\n"
    "       (Cards, Loans, Deposits, Collections, CRM, Notes, Transcripts, Bureau)\n"
    "                                  │\n"
    "                                  ▼\n"
    "                       DATA PRODUCT FACTORY  [Layer 1]\n"
    "           (Ingest raw → Curate silver → Data Contracts → Quality SLA)\n"
    "                                  │\n"
    "                                  ▼\n"
    "                        IDENTITY RESOLUTION\n"
    "           (Deterministic & Probabilistic Matching + Confidence Scoring)\n"
    "                                  │\n"
    "                                  ▼\n"
    "                         GOLDEN FINANCIAL ID\n"
    "           (Conceptual Financial Identity Layer & Reusable Key Anchor)\n"
    "                                  │\n"
    "                                  ▼\n"
    "                           COLLECTIONS 360\n"
    "           (Governed Data Products: Profile, Facilities, Liquidity, Exposure)\n"
    "                                  │\n"
    "                                  ▼\n"
    "                          COLLECTION MEMORY\n"
    "           (Interaction History, Past Promises, Overrides, Transcripts)\n"
    "                                  │\n"
    "                                  ▼\n"
    "                         FINANCIAL ANALYSIS\n"
    "                                  │\n"
    "         ┌────────────────────────┼────────────────────────┐\n"
    "         ▼                        ▼                        ▼\n"
    "   EXPOSURE ANALYSIS       CASH-FLOW STRESS        PAYMENT BEHAVIOUR\n"
    "   (Multi-Product Debt)   (Payroll Disruption)    (PTP Track Record)\n"
    "         └────────────────────────┼────────────────────────┘\n"
    "                                  │\n"
    "                                  ▼\n"
    "                        FINANCIAL HEALTH ENGINE\n"
    "           (Can current capacity support existing obligations?)\n"
    "                                  │\n"
    "                                  ▼\n"
    "                      DECISION INTELLIGENCE  [Layer 3]\n"
    "           (Model Factory · Feature Store · Policy Rules · Hardship)\n"
    "                                  │\n"
    "                                  ▼\n"
    "                            NEXT BEST ACTION\n"
    "                                  │\n"
    "         ┌────────────────────────┼────────────────────────┐\n"
    "         ▼                        ▼                        ▼\n"
    "   DIGITAL REMINDER        SMART PAYMENT PLAN      SPECIALIST OUTREACH\n"
    "  (SMS / WhatsApp)         (Affordability-Driven)  (Empathetic Agent)\n"
    "         └────────────────────────┼────────────────────────┘\n"
    "                                  │\n"
    "                                  ▼\n"
    "                          HUMAN REVIEW & OVERRIDE\n"
    "           (Agent inspects context, approves or modifies action)\n"
    "                                  │\n"
    "                                  ▼\n"
    "                          CUSTOMER INTERACTION\n"
    "           (Appropriate, explainable, and sustainable engagement)\n"
    "                                  │\n"
    "                                  ▼\n"
    "                        OUTCOME & OBSERVATION\n"
    "           (Promise kept, partial payment, hardship reason noted)\n"
    "                                  │\n"
    "                                  ▼\n"
    "                     INCREMENTAL FEEDBACK LOOP\n"
    "           (Updates Collection Memory & C360 Delta in Real Time)\n"
    "                                  ↺")

add_divider(doc)

# ═══════════════════════════════════════════════════════════════════════════════
# SECTION 6 — 3-LAYER HACKATHON ALIGNMENT
# ═══════════════════════════════════════════════════════════════════════════════

add_heading(doc, '6.  Alignment with the Hackathon Three-Layer Challenge', level=1)

add_code_block(doc,
    " ┌──────────────────────────────────────────────────────────────────────┐\n"
    " │                  CROSS-CUTTING GOVERNANCE LAYER                      │\n"
    " │  Data Contracts · Quality SLA · Lineage · RBAC/ABAC · Fair Lending   │\n"
    " │  Protected Attribute Exclusion · Human-in-the-Loop · Audit Trails    │\n"
    " └──────────────────────────────┬───────────────────────────────────────┘\n"
    "                                │\n"
    "       ┌────────────────────────┼────────────────────────┐\n"
    "       │                        │                        │\n"
    "       ▼                        ▼                        ▼\n"
    " ┌───────────────┐  ┌───────────────────┐  ┌───────────────────┐\n"
    " │   LAYER 1     │  │     LAYER 2       │  │     LAYER 3       │\n"
    " │ DATA PRODUCT  │  │  INSIGHT & NLP    │  │  AI DECISIONING   │\n"
    " │ FACTORY       │  │  PLATFORM         │  │  ENGINE           │\n"
    " ├───────────────┤  ├───────────────────┤  ├───────────────────┤\n"
    " │ 9-Source      │  │ Semantic Query    │  │ Next Best Action  │\n"
    " │ Ingestion     │  │ Layer             │  │                   │\n"
    " │ Entity        │  │ NL-to-SQL Engine  │  │ Agent Assist      │\n"
    " │ Resolution    │  │ (SQL shown)       │  │ Cockpit           │\n"
    " │ Golden C360   │  │ RAG over Notes    │  │ Smart Payment     │\n"
    " │ Golden Fin ID │  │ & Transcripts     │  │ Plans             │\n"
    " │ Data          │  │ Collection Memory │  │ Post-Call         │\n"
    " │ Contracts     │  │ Search            │  │ Summaries & QA    │\n"
    " │ Quality SLA   │  │ Out-of-bounds     │  │ Feature Store     │\n"
    " │               │  │ Refusal           │  │ (ML)              │\n"
    " └───────────────┘  └───────────────────┘  └───────────────────┘")

add_divider(doc)

# ═══════════════════════════════════════════════════════════════════════════════
# SECTION 7 — COMPARATIVE TABLE
# ═══════════════════════════════════════════════════════════════════════════════

add_heading(doc, '7.  Comparative Analysis: Automation vs. Intelligence', level=1)

comp_data = [
    ('Operational Philosophy', 'Maximise short-term cash extraction.', 'Maximise sustainable recovery; preserve customer lifetime value.'),
    ('Identity Model', 'Fragmented account IDs siloed by card, loan, or branch.', 'Unified, persistent Golden Financial ID identity layer.'),
    ('Contextual Memory', 'Zero cross-cycle memory; agents start from scratch.', 'Collection Memory: historical attempts, outcomes, transcripts.'),
    ('Financial Assessment', 'Single-account DPD bucket.', 'Multi-dimensional Financial Health: Exposure + Cash Flow + Behaviour.'),
    ('Hardship Treatment', 'Standardised threats, legal notices, aggressive dunning.', 'Fair Collections Intelligence detecting real distress, tailored plans.'),
    ('Payment Plans', 'Rigid settlement matrices with upfront lump sums.', 'Smart Payment Plans based on verified affordability and capacity.'),
    ('AI Role', 'High-velocity robocalls and blind automated SMS.', 'AI decision-support with mandatory human-in-the-loop review.'),
    ('Data Governance', 'Uncontrolled extract files, manual spreadsheets.', 'Formal Data Products, Data Contracts, and full lineage tracking.'),
    ('Customer Experience', 'Harassed, alienated, driven into evasive default.', 'Empathetically treated with transparent rehabilitation pathways.'),
    ('Regulatory Standing', 'High risk of fair-lending regulatory penalties.', 'Fully compliant with ethical recovery, privacy, and fairness mandates.'),
]

tbl2 = doc.add_table(rows=len(comp_data)+1, cols=3)
tbl2.style = 'Table Grid'
headers2 = ['Dimension', 'Legacy Automation', 'Our Customer-Centric Intelligence']
for i, h in enumerate(headers2):
    cell = tbl2.cell(0, i)
    shade_cell(cell, '1E3A5F')
    run = cell.paragraphs[0].add_run(h)
    run.bold = True
    run.font.size = Pt(9)
    run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)

for row_idx, (dim, legacy, ours) in enumerate(comp_data, start=1):
    shade_cell(tbl2.cell(row_idx, 0), 'EFF6FF')
    shade_col = 'FEF2F2' if row_idx % 2 == 0 else 'FFF7ED'
    shade_cell(tbl2.cell(row_idx, 1), shade_col)
    shade_cell(tbl2.cell(row_idx, 2), 'F0FDF4')

    dr = tbl2.cell(row_idx, 0).paragraphs[0].add_run(dim)
    dr.bold = True
    dr.font.size = Pt(9)
    dr.font.color.rgb = RGBColor(0x1E, 0x3A, 0x5F)

    lr = tbl2.cell(row_idx, 1).paragraphs[0].add_run(legacy)
    lr.font.size = Pt(9)
    lr.font.color.rgb = RGBColor(0x6B, 0x21, 0x21)

    or_ = tbl2.cell(row_idx, 2).paragraphs[0].add_run(ours)
    or_.font.size = Pt(9)
    or_.font.color.rgb = RGBColor(0x14, 0x53, 0x2D)

doc.add_paragraph().paragraph_format.space_after = Pt(6)
add_divider(doc)

# ═══════════════════════════════════════════════════════════════════════════════
# SECTION 8 — CONCLUSION
# ═══════════════════════════════════════════════════════════════════════════════

add_heading(doc, '8.  The Ultimate Conceptual Conclusion', level=1)

add_para(doc,
    'Traditional collections platforms treat borrowers as delinquent account numbers to be pressured. '
    'Our platform proves that empathy and data intelligence are not mutually exclusive — they are '
    'mutually reinforcing.', color='374151')

add_callout(doc,
    '"We do not want the collection agent to ask, \'What does this loan record tell me?\'\n\n'
    'We want the system to help the agent ask:\n\n'
    '\'What does the customer\'s complete relevant financial context tell me, what has happened before, '
    'what is changing now, and what is the most appropriate next step?\'"',
    bg='EFF6FF', bar='2563EB', label='CLOSING VISION')

add_code_block(doc,
    "                              THE EVOLUTION\n\n"
    "  FROM:  Data Silos ──→  Collections Case ──→  Aggressive Recovery\n\n"
    "  TO:    Golden Financial ID\n"
    "             ──→  Customer 360\n"
    "                      ──→  Collection Memory\n"
    "                               ──→  Financial Health\n"
    "                                        ──→  Fair Decision Intelligence\n"
    "                                                 ──→  Human-Assisted Action\n"
    "                                                          ──→  Customer Outcome\n"
    "                                                                   ──→  Learning  ↺")

add_para(doc,
    'Unique Value Proposition:',
    bold=True, color='1E3A5F', size=12, space_before=10)

add_code_block(doc,
    "  GOLDEN ID  +  CUSTOMER 360  +  COLLECTION MEMORY\n"
    "          +  FINANCIAL HEALTH  +  SMART PAYMENT PLANNING\n"
    "          +  FAIR COLLECTIONS  +  NEXT BEST ACTION\n\n"
    "  ──────────────────────────────────────────────────\n\n"
    "  IDENTITY  →  MEMORY  →  CONTEXT  →  FINANCIAL UNDERSTANDING\n"
    "           →  DECISION  →  HUMAN ACTION  →  CUSTOMER OUTCOME")

# ─── save ───────────────────────────────────────────────────────────────────

output_path = r'C:\Users\dadid\.gemini\antigravity\scratch\collections-ai\Collections_Intelligence_Platform_Architecture.docx'
doc.save(output_path)
print(f'Document saved -> {output_path}')
